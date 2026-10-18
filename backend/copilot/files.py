#
# Reading the clinic's guidelines out of the files it sends: PDF, Word, text.
#
# Only the text, with small standard libraries (pypdf, python-docx). A scanned
# PDF has no text to extract, and OCR is a dependency and a source of errors
# this does not need yet, so such a file is refused with a sentence saying so,
# and the person can paste its text instead.
#
# What is read is kept as a source: { name, kind, pages, text }. The pasted
# notes are one more source, named "your notes". Every rule Prosper AI writes
# quotes one of them, so the text is kept with the pull request it built.
#

import io
from typing import Optional

NOTES = "your notes"
KINDS = {"pdf": "PDF", "docx": "DOCX", "txt": "TXT", "md": "MD"}
MAX_BYTES = 10 * 1024 * 1024


class Unreadable(ValueError):
    """A file that cannot be read as guidelines. The message says which and why."""


def kind_of(name: str) -> Optional[str]:
    return KINDS.get(name.rsplit(".", 1)[-1].lower()) if "." in name else None


def read(name: str, data: bytes) -> dict:
    """The file as a source, { name, kind, pages, text }. Raises Unreadable."""
    kind = kind_of(name)
    if kind is None:
        raise Unreadable(f"'{name}' is not a PDF, a Word document or a text file.")
    if len(data) > MAX_BYTES:
        raise Unreadable(f"'{name}' is over 10 MB.")
    pages = None
    try:
        if kind == "PDF":
            text, pages = _pdf(data)
        elif kind == "DOCX":
            text = _docx(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except Unreadable:
        raise
    except Exception as error:  # noqa: BLE001 - a broken file is the person's to fix, not a 500
        raise Unreadable(f"'{name}' could not be read: {error}") from error
    text = _tidy(text)
    if not text:
        hint = " It may be a scan: paste its text instead." if kind == "PDF" else ""
        raise Unreadable(f"'{name}' has no text to read.{hint}")
    return {"name": name, "kind": kind, "pages": pages, "text": text}


def notes(text: str) -> Optional[dict]:
    """The pasted notes as a source, or None if there are none."""
    text = _tidy(text)
    return {"name": NOTES, "kind": "NOTES", "pages": None, "text": text} if text else None


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    # Opening hours and price lists are often tables: one line per row.
    for table in document.tables:
        for row in table.rows:
            lines.append(" · ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(lines)


def _tidy(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\r", "").split("\n")]
    # At most one empty line in a row: page breaks and spacing say nothing.
    kept: list[str] = []
    for line in lines:
        if line or (kept and kept[-1]):
            kept.append(line)
    return "\n".join(kept).strip()


def summary(sources: list[dict]) -> str:
    """What was read, as the first step says it: "Read 2 files and your notes · 6 pages"."""
    files = [source for source in sources if source["name"] != NOTES]
    has_notes = len(files) < len(sources)
    pages = sum(source["pages"] or 0 for source in files)
    if not files:
        return "Read your notes"
    read = f"Read {len(files)} {'file' if len(files) == 1 else 'files'}"
    if has_notes:
        read += " and your notes"
    if pages:
        read += f" · {pages} {'page' if pages == 1 else 'pages'}"
    return read
