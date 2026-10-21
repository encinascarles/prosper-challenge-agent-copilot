#
# Prosper AI's pieces, without a model: reading files, turning what the model
# drafts into an agent, the checks every answer must pass, and the retry that
# sends a check's error back. The model is scripted (tests/fakes.py).
#

import io

import pytest

from copilot import artifacts, files
from copilot.llm import Failed
from tests.fakes import ScriptedModel

DRAFT = {
    "persona": "You are the scheduling assistant of Riverside Family Clinic.",
    "initial_node": "greeting",
    "nodes": [
        {
            "name": "greeting",
            "task": "Greet the caller and ask what they need.",
            "end": False,
            "edges": [
                {
                    "description": "Once the caller says what they want.",
                    "target": "Collect details",
                    "fields": [
                        {
                            "name": "Intent",
                            "description": "What they want.",
                            "type": "choice",
                            "options": ["book", "cancel"],
                            "required": True,
                        }
                    ],
                }
            ],
        },
        {
            "name": "Collect details",
            "task": "Ask for the caller's full name.",
            "end": False,
            "edges": [
                {
                    "description": "Once the name is known.",
                    "target": "confirm",
                    "fields": [
                        {
                            "name": "full name",
                            "description": "The caller's full name.",
                            "type": "string",
                            "options": [],
                            "required": False,
                        }
                    ],
                }
            ],
        },
        {"name": "confirm", "task": "Confirm and say goodbye.", "end": True, "edges": []},
    ],
}

SOURCES = [
    {"name": "your notes", "kind": "NOTES", "pages": None, "text": "Never give medical advice."},
    {"name": "hours.txt", "kind": "TXT", "pages": None, "text": "New bookings: Tuesday 10 AM."},
]


def _config() -> dict:
    return artifacts.to_config("Riverside", DRAFT)


# ---- the agent ---------------------------------------------------------------


def test_a_draft_becomes_an_agent_the_builder_accepts():
    config = _config()

    artifacts.check_agent(config)
    greeting, details, confirm = config["nodes"]
    assert [node["name"] for node in config["nodes"]] == ["greeting", "collect_details", "confirm"]
    assert config["model"] == "gpt-4o"
    assert greeting["task_messages"] == [
        {"role": "developer", "content": "Greet the caller and ask what they need."}
    ]
    (edge,) = greeting["edges"]
    assert edge["function"] == "go_to_collect_details"
    assert edge["target"] == "collect_details"
    assert edge["properties"] == {
        "intent": {"type": "string", "enum": ["book", "cancel"], "description": "What they want."}
    }
    assert edge["required"] == ["intent"]
    assert details["edges"][0]["required"] == []
    assert confirm == {"name": "confirm", "task_messages": confirm["task_messages"], "end": True}


def test_two_edges_to_one_step_get_numbered_functions():
    draft = {**DRAFT, "nodes": [dict(node) for node in DRAFT["nodes"]]}
    edge = DRAFT["nodes"][0]["edges"][0]
    draft["nodes"][0] = {**DRAFT["nodes"][0], "edges": [edge, {**edge, "description": "Again."}]}

    functions = [e["function"] for e in artifacts.to_config("R", draft)["nodes"][0]["edges"]]

    assert functions == ["go_to_collect_details", "go_to_collect_details_2"]


def test_an_agent_and_its_draft_go_both_ways():
    config = _config()

    assert artifacts.to_config("Riverside", artifacts.from_config(config)) == config


def test_an_unknown_target_is_refused_as_a_save_would_say_it():
    config = _config()
    config["nodes"][1]["edges"][0]["target"] = "goodbye"

    with pytest.raises(ValueError, match="targets unknown node 'goodbye'"):
        artifacts.check_agent(config)


def test_nothing_may_lead_back_into_the_start_step():
    config = _config()
    config["nodes"][1]["edges"][0]["target"] = "greeting"

    with pytest.raises(ValueError, match="goes back to the start step 'greeting'"):
        artifacts.check_agent(config)


@pytest.mark.parametrize(
    "change, problem",
    [
        (lambda c: c["nodes"][2].pop("end"), "No step ends the call"),
        (lambda c: c["nodes"][1].pop("edges"), "'collect_details' has no edges and does not end"),
        (lambda c: c["nodes"][2].update(edges=c["nodes"][0]["edges"]), "cannot have edges"),
    ],
)
def test_a_flow_a_call_cannot_get_through_is_refused(change, problem):
    config = _config()
    change(config)

    with pytest.raises(ValueError, match=problem):
        artifacts.check_agent(config)


def test_a_change_keeps_what_the_draft_has_no_place_for():
    before = _config()
    before["nodes"][0]["role_message"] = "Speak Spanish."
    before["voice_id"] = "another-voice"

    config = artifacts.keep_extras(artifacts.to_config("Riverside", DRAFT), before)

    assert config["nodes"][0]["role_message"] == "Speak Spanish."
    assert config["voice_id"] == "another-voice"


# ---- rules, tests, questions -------------------------------------------------


def _rule(**changes) -> dict:
    rule = {
        "step": "Collect details",
        "text": "No medical advice.",
        "source": "your notes",
        "quote": "never give  Medical advice",
    }
    return {**rule, **changes}


def _check(rule: dict, answers: tuple = ()) -> tuple:
    steps = {node["name"] for node in _config()["nodes"]}
    return artifacts.check_rule(rule, steps), artifacts.check_quote(rule, SOURCES, list(answers))


def test_a_rule_quotes_its_source_and_names_a_step():
    step, source = _check(_rule())

    assert step == "collect_details"
    assert source == {
        "kind": "guidelines",
        "file": "your notes",
        "quote": "never give  Medical advice",
    }
    # Or one of the clinic's answers, as given.
    answer = _rule(source="your answer", quote="Offer the waitlist")
    assert _check(answer, ("Offer the waitlist",))[1] == {
        "kind": "answer",
        "quote": "Offer the waitlist",
    }
    with pytest.raises(ValueError, match="not one of the clinic's answers"):
        _check(answer, ("Ask them to call back",))


@pytest.mark.parametrize(
    "changes, problem",
    [
        ({"quote": "Always give medical advice."}, "is not in 'your notes' word for word"),
        ({"source": "rules.pdf"}, "not one of the sources: your notes, hours.txt, your answer"),
        ({"step": "waitlist"}, "is for step 'waitlist', which the agent does not have"),
    ],
)
def test_a_rule_that_cannot_be_traced_is_refused(changes, problem):
    with pytest.raises(ValueError, match=problem):
        _check(_rule(**changes))


def test_a_path_check_must_name_a_step_or_a_field():
    good = {"name": "Books", "caller": "Maria.", "checks": ["Reaches Confirm", "Says hello"]}
    recorded = {**good, "checks": ["Full name is recorded", "Ends the call"]}
    artifacts.check_tests([good, recorded], _config())

    with pytest.raises(ValueError, match="'Reaches Waitlist' names no step"):
        artifacts.check_tests(
            [{**good, "checks": ["Reaches Waitlist", "Ends the call"]}], _config()
        )
    with pytest.raises(ValueError, match="'Date is recorded' names no field"):
        artifacts.check_tests(
            [{**good, "checks": ["Date is recorded", "Ends the call"]}], _config()
        )


@pytest.mark.parametrize(
    "checks, problem",
    [
        (["Ends the call"], "has 1 checks: give it 2 to 4"),
        (["Says hello", "Reads back the time"], "has 2 judged checks"),
        (["Ends the call", "Keeps every reply to one or two sentences"], "checks style or wording"),
        (["Ends the call", "Speaks in a warm tone"], "checks style or wording"),
    ],
)
def test_a_test_checks_what_happens_not_how_it_is_said(checks, problem):
    test = {"name": "Books", "caller": "Maria.", "checks": checks}

    with pytest.raises(ValueError, match=problem):
        artifacts.check_tests([test], _config())


def test_a_task_is_three_short_sentences_and_a_rule_one():
    config = _config()
    long = "Greet them. Ask their name. Ask why they call. Offer to help."
    config["nodes"][0]["task_messages"][0]["content"] = long

    with pytest.raises(ValueError, match="task of step 'greeting' is 4 sentences"):
        artifacts.check_concise(config)
    # A change is held to it only where it changes something.
    artifacts.check_concise(config, before=config)
    with pytest.raises(ValueError, match="must be one short sentence"):
        artifacts.check_rule_text("Be kind. Be quick.")


@pytest.mark.parametrize(
    "rule",
    [
        "Ask whether it is their first visit before taking any outgoing edge.",
        "Never skip the confirm node.",
        "Call the go_to_goodbye function only once they are done.",
        "Make the transition to Nurse callback when symptoms come up.",
    ],
)
def test_a_rule_says_steps_and_moving_on_never_the_graph(rule):
    with pytest.raises(ValueError, match="say it in the product's words"):
        artifacts.check_rule_text(rule)


def test_a_task_or_condition_that_changes_is_held_to_the_product_words():
    config = _config()
    config["nodes"][0]["task_messages"][0]["content"] = "Greet them, then take the edge."
    with pytest.raises(ValueError, match="task of step 'greeting' says \"edge\""):
        artifacts.check_concise(config)
    config = _config()
    config["nodes"][0]["edges"][0]["description"] = "Take this edge once they said why."
    with pytest.raises(ValueError, match="A condition of step 'greeting' says \"edge\""):
        artifacts.check_concise(config)
    # What it does not change is left as it is.
    artifacts.check_concise(config, before=config)
    # The product's own words pass.
    artifacts.check_rule_text("Ask whether it is their first visit before moving on.")


def test_questions_are_few_with_two_or_three_options():
    steps = {node["name"] for node in _config()["nodes"]}
    option = {"label": "Waitlist", "rule": {"step": None, "text": "Offer the waitlist."}}
    question = {"text": "What if no time works?", "options": [option, option]}
    artifacts.check_questions([question, question], steps)

    with pytest.raises(ValueError, match="2 questions at most"):
        artifacts.check_questions([question] * 3, steps)
    with pytest.raises(ValueError, match="needs 2 or 3 options"):
        artifacts.check_questions([{**question, "options": [option]}], steps)
    with pytest.raises(ValueError, match="for step 'Waitlist', which the agent does not have"):
        wrong = {**option, "rule": {"step": "Waitlist", "text": "Offer the waitlist."}}
        artifacts.check_questions([{**question, "options": [wrong, option]}], steps)
    # Before anything is drawn: three at most, and an option is only a choice.
    picked = {"label": "Waitlist", "rule": None}
    artifacts.check_questions([{**question, "options": [picked, picked]}] * 3, None)
    with pytest.raises(ValueError, match="Nothing is drawn yet"):
        artifacts.check_questions([question], None)


# ---- the retry ---------------------------------------------------------------


def test_a_refused_answer_goes_back_with_its_error():
    broken = {**DRAFT, "initial_node": "start"}
    model = ScriptedModel({"title": [broken, DRAFT]})

    def check(answer):
        artifacts.check_agent(artifacts.to_config("R", answer))

    assert model.ask("title", "Guidelines.", artifacts.AGENT, check) == DRAFT
    first, second = model.asked["title"]
    assert len(first) == 2
    assert "initial_node 'start' is not a defined node" in second[-1]["content"]
    assert second[-2]["role"] == "assistant"


def test_it_gives_up_after_three_refusals():
    model = ScriptedModel({"title": [{**DRAFT, "initial_node": "start"}] * 3})

    with pytest.raises(Failed, match="initial_node 'start'"):
        model.ask(
            "title",
            "x",
            artifacts.AGENT,
            lambda a: artifacts.check_agent(artifacts.to_config("R", a)),
        )
    assert len(model.asked["title"]) == 3


def test_an_answer_cut_off_is_asked_for_again_once():
    from copilot.llm import CutOff

    class Stuck(ScriptedModel):
        def _complete(self, messages, name, schema):
            self.asked.setdefault(name, []).append(list(messages))
            raise CutOff("The model ran out of room before it finished its answer.")

    model = Stuck({})
    with pytest.raises(CutOff):
        model.ask("title", "x", artifacts.AGENT)
    assert len(model.asked["title"]) == 2
    # The same request again, not the cut-off answer fed back.
    assert model.asked["title"][0] == model.asked["title"][1]


def test_usage_is_priced_by_model():
    from types import SimpleNamespace

    model = ScriptedModel({})
    model.model = "gpt-6.1-sol"
    details = SimpleNamespace(cached_tokens=1000)
    usage = SimpleNamespace(
        prompt_tokens=11000, completion_tokens=2000, prompt_tokens_details=details
    )

    model.count(usage)

    # 10k fresh at $2, 1k cached at $0.10, 2k out at $10, per million.
    assert model.usage["usd"] == pytest.approx(0.0401)
    assert (model.usage["prompt"], model.usage["cached"], model.usage["completion"]) == (
        11000,
        1000,
        2000,
    )


# ---- files -------------------------------------------------------------------


def test_text_and_word_files_are_read():
    import docx

    document = docx.Document()
    document.add_paragraph("Opening hours")
    document.add_paragraph("Monday to   Friday, 8 AM to 6 PM.")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Saturday"
    data = io.BytesIO()
    document.save(data)

    read = files.read("Hours.docx", data.getvalue())

    assert read["kind"] == "DOCX"
    assert read["text"] == "Opening hours\nMonday to Friday, 8 AM to 6 PM.\nSaturday"
    assert files.read("notes.md", b"# Hi\n\n\n\nThere")["text"] == "# Hi\n\nThere"


def test_a_pdf_without_text_is_refused_with_why():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    data = io.BytesIO()
    writer.write(data)

    with pytest.raises(files.Unreadable, match="'scan.pdf' has no text to read. It may be a scan"):
        files.read("scan.pdf", data.getvalue())


def test_a_file_of_another_kind_is_refused():
    with pytest.raises(files.Unreadable, match="not a PDF, a Word document or a text file"):
        files.read("photo.png", b"...")


def test_what_was_read_is_summed_up():
    pdf = {"name": "a.pdf", "kind": "PDF", "pages": 4, "text": "x"}
    word = {"name": "b.docx", "kind": "DOCX", "pages": None, "text": "x"}
    notes = files.notes("  Some notes  ")

    assert files.summary([notes, pdf, word]) == "Read 2 files and your notes · 4 pages"
    assert files.summary([pdf]) == "Read 1 file · 4 pages"
    assert files.summary([notes]) == "Read your notes"
    assert files.notes("   ") is None


def test_every_problem_of_an_answer_goes_back_at_once():
    one = {"name": "A", "caller": "x", "checks": ["Ends the call"]}
    two = {"name": "B", "caller": "x", "checks": ["Ends the call", "Speaks in a warm tone"]}

    with pytest.raises(ValueError) as refused:
        artifacts.check_tests([one, two], _config())

    assert str(refused.value).startswith("2 problems:")
    assert "Test 'A' has 1 checks" in str(refused.value)
    assert "checks style or wording" in str(refused.value)
